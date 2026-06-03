#!/usr/bin/env python3
"""
JM NFE — Extrator de dados de liquidação a partir do PDF do processo.

O Agente lê o PDF inteiro (sentença/acórdão/peças) e extrai, de forma ESTRUTURADA,
todos os dados necessários para montar a liquidação trabalhista. Usa a Claude API
(claude-opus-4-8) com saída estruturada (structured output) validada por schema.

PRINCÍPIO PERICIAL (inegociável): a IA PROPÕE, o PERITO DECIDE.
- Só extrai o que está no documento; nunca inventa valor, índice ou data.
- Cita o trecho-fonte de cada dado e marca o que ficou incerto (revisar).
- Índices de correção/juros NÃO são fixados pela IA: o critério segue a decisão
  e o período, e deve ser confirmado na fonte oficial pelo perito.

Uso:
    export ANTHROPIC_API_KEY="sua-chave"
    python extrator_liquidacao.py processo.pdf
    python extrator_liquidacao.py processo.pdf --out ./saida --planilha ../07-planilha-liquidacao-avancada.xlsx
    python extrator_liquidacao.py processo.pdf --pdf-nativo     # PDF escaneado (manda o PDF p/ o modelo)

Saídas:
    <out>/<nome>_extracao.json   -> dados estruturados
    <out>/<nome>_relatorio.md    -> relatório legível para conferência do perito
    (opcional) cópia da planilha 07 com Verbas e Memória pré-preenchidas

Dependências: anthropic, pdfplumber, pydantic, openpyxl
    pip install anthropic pdfplumber pydantic openpyxl
"""
import argparse
import base64
import json
import os
import sys
from pathlib import Path
from typing import List, Optional, Literal

try:
    from pydantic import BaseModel, Field
except ImportError:
    sys.exit("Falta a dependência 'pydantic'. Rode: pip install pydantic")

MODELO_PADRAO = "claude-opus-4-8"

# ----------------------------------------------------------------------------
# 1) SCHEMA DE SAÍDA (o que precisamos para a liquidação)
# ----------------------------------------------------------------------------
Confianca = Literal["alta", "media", "baixa"]


class Verba(BaseModel):
    descricao: str = Field(description="Nome da verba deferida (ex.: 'Horas extras 50%').")
    periodo_inicio: Optional[str] = Field(None, description="Início do período (MM/AAAA), se houver.")
    periodo_fim: Optional[str] = Field(None, description="Fim do período (MM/AAAA), se houver.")
    reflexos: Optional[str] = Field(None, description="Reflexos deferidos (DSR, 13º, férias+1/3, FGTS...).")
    base_calculo: Optional[str] = Field(None, description="Base de cálculo, se especificada.")
    observacao: Optional[str] = Field(None, description="Observações relevantes (percentuais, limites).")
    confianca: Confianca = Field(description="Confiança da extração desta verba.")


class FonteTrecho(BaseModel):
    campo: str = Field(description="Nome do campo extraído (ex.: 'data_admissao').")
    trecho: str = Field(description="Trecho LITERAL do documento que fundamenta o dado (curto).")


class DadosLiquidacao(BaseModel):
    # Identificação
    numero_processo: Optional[str] = None
    vara: Optional[str] = None
    comarca: Optional[str] = None
    reclamante: Optional[str] = None
    reclamada: Optional[str] = None
    # Contrato
    data_admissao: Optional[str] = Field(None, description="DD/MM/AAAA.")
    data_afastamento: Optional[str] = Field(None, description="DD/MM/AAAA (rescisão/afastamento).")
    motivo_saida: Optional[str] = Field(None, description="Ex.: dispensa sem justa causa, pedido, rescisão indireta.")
    ultimo_salario: Optional[str] = Field(None, description="Último salário-base (R$), como consta.")
    funcao: Optional[str] = None
    jornada: Optional[str] = None
    cct_aplicavel: Optional[str] = Field(None, description="Convenção/Acordo coletivo citado, se houver.")
    # Processuais
    data_ajuizamento: Optional[str] = Field(None, description="DD/MM/AAAA — marca a virada pré-judicial->judicial.")
    criterio_correcao: Optional[str] = Field(None, description="Critério de correção monetária fixado na decisão, se houver.")
    criterio_juros: Optional[str] = Field(None, description="Critério de juros fixado na decisão, se houver.")
    criterio_definido_na_sentenca: Optional[bool] = Field(None, description="A decisão fixou expressamente o critério de atualização?")
    # FGTS
    fgts_informacoes: Optional[str] = Field(None, description="Dados de FGTS citados (depósitos, recolhimento).")
    multa_40: Optional[bool] = Field(None, description="A multa de 40% do FGTS é devida, conforme o documento?")
    # Verbas deferidas
    verbas_deferidas: List[Verba] = Field(default_factory=list)
    # Meta / rastreabilidade
    fontes: List[FonteTrecho] = Field(default_factory=list, description="Trechos-fonte dos principais campos.")
    campos_para_revisar: List[str] = Field(default_factory=list, description="Campos incertos/ausentes que o perito deve conferir.")
    observacoes_extracao: str = Field("", description="Notas da extração: ambiguidades, ressalvas, o que não foi encontrado.")


# ----------------------------------------------------------------------------
# 2) INSTRUÇÕES DO AGENTE (system prompt) — método pericial + guardrails
# ----------------------------------------------------------------------------
SYSTEM_PROMPT = """\
Você é um assistente técnico de perícia trabalhista da JM NFE Consultoria. Sua tarefa é
LER o documento de um processo (sentença, acórdão, peças, documentos anexos) e EXTRAIR,
de forma estruturada, os dados necessários para a liquidação trabalhista.

REGRAS INEGOCIÁVEIS (disciplina pericial):
1. ZERO INVENÇÃO. Extraia somente o que está EXPLÍCITO no documento. Se um dado não
   aparece, deixe o campo nulo e acrescente o nome do campo em "campos_para_revisar".
2. CITE A FONTE. Para os campos principais (datas, salário, verbas, critério de
   atualização, ajuizamento), inclua em "fontes" o trecho LITERAL e curto que fundamenta
   o dado. Sem trecho de apoio, marque confiança "baixa".
3. NÃO FIXE ÍNDICES. Você não decide índice de correção/juros. Se a decisão fixar um
   critério, transcreva-o em "criterio_correcao"/"criterio_juros" e marque
   "criterio_definido_na_sentenca" = true. Caso contrário, deixe nulo, marque false e
   inclua "criterio_correcao" e "criterio_juros" em "campos_para_revisar".
4. PERÍODOS DAS VERBAS. Para cada verba deferida, capture o período (MM/AAAA) e os
   reflexos exatamente como constam. Não generalize de uma verba para outra.
5. INCERTEZA EXPLÍCITA. Use confiança "alta" só quando o texto for inequívoco; "media"
   quando exigir interpretação; "baixa" quando for indício frágil.
6. VOCÊ NÃO ADVOGA E NÃO CALCULA AQUI. Apenas extrai dados para o perito conferir e
   calcular. Não emita tese jurídica nem valor de liquidação.
7. LGPD: trate os dados pessoais com sigilo; não acrescente informação externa ao documento.

Responda preenchendo o schema estruturado solicitado. Seja fiel ao documento e conservador.
"""

USER_INSTRUCAO = (
    "Extraia os dados de liquidação do documento abaixo, seguindo as regras. "
    "Cite trechos-fonte e liste o que precisa de revisão.\n\n"
    "=== DOCUMENTO DO PROCESSO ===\n"
)


# ----------------------------------------------------------------------------
# 3) LEITURA DO PDF
# ----------------------------------------------------------------------------
def extrair_texto_pdf(caminho: Path) -> str:
    """Extrai texto do PDF, página a página, com marcadores de página."""
    try:
        import pdfplumber
    except ImportError:
        sys.exit("Falta a dependência 'pdfplumber'. Rode: pip install pdfplumber")

    partes: List[str] = []
    with pdfplumber.open(str(caminho)) as pdf:
        total = len(pdf.pages)
        for i, page in enumerate(pdf.pages, start=1):
            txt = page.extract_text() or ""
            partes.append(f"\n[Página {i}/{total}]\n{txt}")
    return "".join(partes).strip()


def pdf_como_base64(caminho: Path) -> str:
    return base64.standard_b64encode(caminho.read_bytes()).decode("utf-8")


# ----------------------------------------------------------------------------
# 4) CHAMADA À CLAUDE API (structured output)
# ----------------------------------------------------------------------------
def extrair_com_claude(conteudo_usuario, modelo: str) -> tuple[DadosLiquidacao, dict]:
    """conteudo_usuario: str (texto) OU lista de blocos (PDF nativo)."""
    try:
        import anthropic
    except ImportError:
        sys.exit("Falta a dependência 'anthropic'. Rode: pip install anthropic")

    if not (os.getenv("ANTHROPIC_API_KEY") or os.getenv("ANTHROPIC_AUTH_TOKEN")):
        sys.exit("Defina ANTHROPIC_API_KEY no ambiente antes de rodar.")

    client = anthropic.Anthropic()

    # System como bloco com cache_control: as instruções são estáveis entre execuções,
    # então o prefixo pode ser servido do cache (economia em lotes de processos).
    system = [{"type": "text", "text": SYSTEM_PROMPT, "cache_control": {"type": "ephemeral"}}]

    if isinstance(conteudo_usuario, str):
        messages = [{"role": "user", "content": USER_INSTRUCAO + conteudo_usuario}]
    else:
        messages = [{"role": "user", "content": conteudo_usuario}]

    resp = client.messages.parse(
        model=modelo,
        max_tokens=16000,
        thinking={"type": "adaptive"},   # deixa o modelo raciocinar sobre o texto jurídico
        system=system,
        messages=messages,
        output_format=DadosLiquidacao,
    )

    if resp.stop_reason == "refusal":
        sys.exit("O modelo recusou a extração (verifique o conteúdo do documento).")

    dados = resp.parsed_output
    if dados is None:
        sys.exit("Não foi possível obter a saída estruturada (parsed_output vazio).")

    uso = {
        "input_tokens": getattr(resp.usage, "input_tokens", None),
        "output_tokens": getattr(resp.usage, "output_tokens", None),
        "cache_read_input_tokens": getattr(resp.usage, "cache_read_input_tokens", None),
        "cache_creation_input_tokens": getattr(resp.usage, "cache_creation_input_tokens", None),
    }
    return dados, uso


# ----------------------------------------------------------------------------
# 5) RELATÓRIO LEGÍVEL PARA O PERITO
# ----------------------------------------------------------------------------
def gerar_relatorio_md(d: DadosLiquidacao, origem: str, uso: dict) -> str:
    def v(x):
        return x if x not in (None, "") else "_(não encontrado)_"

    linhas = [
        f"# Extração de liquidação — conferência do perito",
        f"> Fonte: `{origem}`  ·  Gerado pelo extrator JM NFE (claude `{MODELO_PADRAO}`).",
        "",
        "> ⚠️ A IA PROPÕE, o PERITO DECIDE. Confira cada campo no documento antes de calcular.",
        "> Índices de correção/juros NÃO são fixados pela IA — seguem a decisão e o período.",
        "",
        "## Identificação",
        f"- Processo: **{v(d.numero_processo)}**  ·  Vara: {v(d.vara)}  ·  Comarca: {v(d.comarca)}",
        f"- Reclamante: {v(d.reclamante)}  ·  Reclamada: {v(d.reclamada)}",
        "",
        "## Contrato",
        f"- Admissão: **{v(d.data_admissao)}**  ·  Afastamento: **{v(d.data_afastamento)}**",
        f"- Motivo da saída: {v(d.motivo_saida)}",
        f"- Último salário: **{v(d.ultimo_salario)}**  ·  Função: {v(d.funcao)}",
        f"- Jornada: {v(d.jornada)}  ·  CCT/ACT: {v(d.cct_aplicavel)}",
        "",
        "## Dados processuais / atualização",
        f"- Data de ajuizamento: **{v(d.data_ajuizamento)}**",
        f"- Critério fixado na decisão? **{('Sim' if d.criterio_definido_na_sentenca else 'Não/Não identificado')}**",
        f"- Correção monetária: {v(d.criterio_correcao)}",
        f"- Juros: {v(d.criterio_juros)}",
        "",
        "## FGTS",
        f"- Informações: {v(d.fgts_informacoes)}",
        f"- Multa de 40%: {('Sim' if d.multa_40 else 'Não/Não identificado')}",
        "",
        "## Verbas deferidas",
    ]
    if d.verbas_deferidas:
        linhas.append("| Verba | Período | Reflexos | Base | Conf. | Obs. |")
        linhas.append("|---|---|---|---|---|---|")
        for vb in d.verbas_deferidas:
            periodo = " a ".join([p for p in (vb.periodo_inicio, vb.periodo_fim) if p]) or "—"
            linhas.append(
                f"| {vb.descricao} | {periodo} | {vb.reflexos or '—'} | "
                f"{vb.base_calculo or '—'} | {vb.confianca} | {vb.observacao or '—'} |"
            )
    else:
        linhas.append("_(nenhuma verba identificada — revisar manualmente)_")

    linhas += ["", "## ⚑ Campos para revisar (conferência obrigatória)"]
    if d.campos_para_revisar:
        linhas += [f"- [ ] {c}" for c in d.campos_para_revisar]
    else:
        linhas.append("- _(nenhum sinalizado — ainda assim, confira os campos-chave)_")

    linhas += ["", "## Trechos-fonte citados"]
    if d.fontes:
        for f in d.fontes:
            trecho = (f.trecho or "").replace("\n", " ").strip()
            linhas.append(f"- **{f.campo}**: “{trecho}”")
    else:
        linhas.append("- _(sem trechos citados — tratar extração com cautela)_")

    if d.observacoes_extracao:
        linhas += ["", "## Observações da extração", d.observacoes_extracao]

    if any(uso.values()):
        linhas += ["", "---", f"_Uso de tokens: {json.dumps(uso, ensure_ascii=False)}_"]

    linhas += [
        "",
        "---",
        "_Material técnico de apoio; não constitui parecer jurídico nem cálculo de "
        "liquidação. Requer conferência do perito e confirmação dos índices na fonte oficial._",
    ]
    return "\n".join(linhas)


# ----------------------------------------------------------------------------
# 6) (OPCIONAL) INJETAR NA PLANILHA DE LIQUIDAÇÃO 07
# ----------------------------------------------------------------------------
def preencher_planilha(d: DadosLiquidacao, modelo_xlsx: Path, destino_xlsx: Path) -> None:
    """Copia a planilha 07 e pré-preenche a aba 'Verbas' e a 'Memoria'."""
    try:
        import openpyxl
    except ImportError:
        print("openpyxl ausente — pulando preenchimento da planilha.", file=sys.stderr)
        return
    if not modelo_xlsx.exists():
        print(f"Planilha modelo não encontrada: {modelo_xlsx}", file=sys.stderr)
        return

    wb = openpyxl.load_workbook(str(modelo_xlsx))
    # Desprotege para escrever (o arquivo de saída é uma cópia de trabalho)
    for ws in wb.worksheets:
        ws.protection.sheet = False

    # Verbas: descrição na coluna B, a partir da linha 5
    if "Verbas" in wb.sheetnames and d.verbas_deferidas:
        wv = wb["Verbas"]
        linha = 5
        for vb in d.verbas_deferidas:
            if linha > 29:  # respeita o range do template
                break
            wv.cell(row=linha, column=2, value=vb.descricao)  # B: parcela
            # C (competência) e D (valor) ficam para o perito preencher por competência
            linha += 1

    # Memoria: registra premissas extraídas (col C = critério, D = fonte/observação)
    if "Memoria" in wb.sheetnames:
        wm = wb["Memoria"]
        mapa = {
            "Período/competências das parcelas":
                (f"{d.data_admissao or '?'} a {d.data_afastamento or '?'}", "extraído do PDF"),
            "Critério de correção (fase)":
                (d.criterio_correcao or "VERIFICAR (não fixado)", "decisão liquidanda"),
            "Critério de juros":
                (d.criterio_juros or "VERIFICAR (não fixado)", "decisão liquidanda"),
        }
        # As premissas começam na linha 5 (col B = nome, C = critério, D = fonte)
        for row in range(5, 15):
            nome = wm.cell(row=row, column=2).value
            if nome in mapa:
                criterio, fonte = mapa[nome]
                wm.cell(row=row, column=3, value=criterio)
                wm.cell(row=row, column=4, value=fonte)

    wb.save(str(destino_xlsx))
    print(f"Planilha pré-preenchida salva em: {destino_xlsx}")


# ----------------------------------------------------------------------------
# 7) CLI
# ----------------------------------------------------------------------------
def main() -> None:
    ap = argparse.ArgumentParser(description="Extrai dados de liquidação de um PDF de processo (JM NFE).")
    ap.add_argument("pdf", type=Path, help="Caminho do PDF do processo.")
    ap.add_argument("--out", type=Path, default=Path("."), help="Pasta de saída (padrão: atual).")
    ap.add_argument("--modelo", default=MODELO_PADRAO, help=f"Modelo Claude (padrão: {MODELO_PADRAO}).")
    ap.add_argument("--pdf-nativo", action="store_true",
                    help="Envia o PDF ao modelo (para PDFs escaneados/sem texto extraível).")
    ap.add_argument("--planilha", type=Path, default=None,
                    help="Planilha 07 modelo (.xlsx) para pré-preencher uma cópia.")
    args = ap.parse_args()

    if not args.pdf.exists():
        sys.exit(f"PDF não encontrado: {args.pdf}")
    args.out.mkdir(parents=True, exist_ok=True)
    base = args.pdf.stem

    # Monta o conteúdo do usuário (texto ou PDF nativo)
    if args.pdf_nativo:
        print("Modo PDF nativo: enviando o PDF ao modelo (escaneado).", file=sys.stderr)
        conteudo = [
            {"type": "document",
             "source": {"type": "base64", "media_type": "application/pdf",
                        "data": pdf_como_base64(args.pdf)}},
            {"type": "text", "text": USER_INSTRUCAO.replace("=== DOCUMENTO DO PROCESSO ===", "")},
        ]
    else:
        texto = extrair_texto_pdf(args.pdf)
        if len(texto) < 200:
            print("AVISO: pouco texto extraído — o PDF pode ser escaneado. "
                  "Reexecute com --pdf-nativo.", file=sys.stderr)
        if len(texto) > 600_000:
            print("AVISO: documento muito grande; considere recortar as peças relevantes.",
                  file=sys.stderr)
        conteudo = texto

    print(f"Extraindo dados com {args.modelo}...", file=sys.stderr)
    dados, uso = extrair_com_claude(conteudo, args.modelo)

    # Salva JSON
    json_path = args.out / f"{base}_extracao.json"
    json_path.write_text(dados.model_dump_json(indent=2), encoding="utf-8")
    # Salva relatório
    md_path = args.out / f"{base}_relatorio.md"
    md_path.write_text(gerar_relatorio_md(dados, args.pdf.name, uso), encoding="utf-8")

    print(f"OK: {json_path}")
    print(f"OK: {md_path}")

    if args.planilha:
        destino = args.out / f"{base}_liquidacao_preenchida.xlsx"
        preencher_planilha(dados, args.planilha, destino)

    # Resumo no terminal
    n = len(dados.verbas_deferidas)
    rev = len(dados.campos_para_revisar)
    print(f"Resumo: {n} verba(s) identificada(s); {rev} campo(s) para revisar.")
    print("Lembrete: a IA PROPÕE, o PERITO DECIDE — confira tudo no documento.")


if __name__ == "__main__":
    main()
